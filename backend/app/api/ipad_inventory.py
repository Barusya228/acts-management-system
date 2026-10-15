import uuid
import csv
import base64
from io import BytesIO, StringIO
from typing import Literal, Optional
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Query, status
from sqlalchemy import BigInteger, String, case, cast, func, or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from app.core.database import get_db
from app.core.deps import get_current_admin_user, get_current_guest_or_admin_user
from app.db.models import Act, IpadActAppendix, IpadDevice, IpadStudentAssignment, User
from app.db.states import ACTIVE_IPAD_ASSIGNMENT_STATUSES
from app.schemas.schemas import IpadAvailableResolveRequest, IpadDeviceBulkCreate, IpadDeviceCreate, IpadDeviceUpdate, IpadInventorySyncRequest
from app.services.audit_service import record_audit


router = APIRouter()


def _serialize(
    device: IpadDevice,
    assignment: IpadStudentAssignment | None = None,
    duplicate_tag_count: int = 1,
    has_history: bool = False,
) -> dict:
    return {
        "id": str(device.id),
        "device_name": device.device_name,
        "model": device.model,
        "tag": device.tag,
        "serial_number": device.serial_number,
        "status": device.status,
        "notes": device.notes,
        "student_name": assignment.student_name if assignment else None,
        "act_id": str(assignment.act_id) if assignment else None,
        "created_at": device.created_at.isoformat(),
        "updated_at": device.updated_at.isoformat(),
        "duplicate_tag_count": duplicate_tag_count,
        "has_history": has_history,
        "is_archived": device.is_archived,
    }


def _payload_device_ids(payload: object) -> set[str]:
    if not isinstance(payload, dict):
        return set()
    device_ids = {
        str(payload.get(key))
        for key in ("device_id", "old_device_id", "new_device_id")
        if payload.get(key)
    }
    items = payload.get("items")
    if isinstance(items, list):
        device_ids.update(
            str(item.get("device_id"))
            for item in items
            if isinstance(item, dict) and item.get("device_id")
        )
    return device_ids


def _history_device_ids(db: Session, device_ids: list[UUID]) -> set[UUID]:
    if not device_ids:
        return set()
    history_ids = {
        item_id
        for item_id, in db.query(IpadStudentAssignment.ipad_device_id).filter(
            IpadStudentAssignment.ipad_device_id.in_(device_ids)
        ).distinct().all()
        if item_id
    }
    string_ids = {str(item_id): item_id for item_id in device_ids}
    appendix_filters = [
        cast(IpadActAppendix.payload_json, String).contains(string_id)
        for string_id in string_ids
    ]
    payloads = db.query(IpadActAppendix.payload_json).filter(or_(*appendix_filters)).all()
    for payload, in payloads:
        for referenced_id in _payload_device_ids(payload):
            if referenced_id in string_ids:
                history_ids.add(string_ids[referenced_id])
    return history_ids


def _parse_ipad_sync_csv(csv_text: str) -> tuple[list[dict], list[str]]:
    parsed: list[dict] = []
    errors: list[str] = []
    reader = csv.reader(StringIO(csv_text.lstrip("\ufeff")))
    for row in reader:
        line_number = reader.line_num
        if not row or not any(str(value).strip() for value in row):
            continue
        # Некоторые MDM-экспорты оборачивают целую CSV-строку в одну колонку:
        # "46303,""Model"",...". Разворачиваем вложенную строку повторно.
        if len(row) == 1 and "," in str(row[0]):
            row = next(csv.reader([str(row[0])]))
        normalized = [str(value).strip() for value in row]
        first_column = normalized[0].casefold() if normalized else ""
        fourth_column = normalized[3].casefold() if len(normalized) > 3 else ""
        if reader.line_num == 1 and (
            first_column in {"tag", "tags", "asset tag", "ipad tag"}
            or fourth_column in {"serial", "serial number", "serial_number"}
        ):
            continue
        if len(normalized) < 4:
            errors.append(f"Строка {line_number}: ожидается минимум 4 колонки")
            continue
        tag, model, _device_name, serial_number = normalized[:4]
        if not tag or not serial_number:
            errors.append(f"Строка {line_number}: Tag и Serial Number обязательны")
            continue
        parsed.append({
            "line_number": line_number,
            "tag": tag,
            "model": model or None,
            "device_name": "iPad",
            "serial_number": serial_number,
            # Назначения оформляются только актами SmartAct. Дополнительные
            # колонки MDM (Asset Tag / user / group) намеренно игнорируются.
            "assigned_user": "",
            "group": "",
        })

    serial_lines: dict[str, list[int]] = {}
    tag_lines: dict[str, list[int]] = {}
    for row in parsed:
        serial_lines.setdefault(row["serial_number"].casefold(), []).append(row["line_number"])
        tag_lines.setdefault(row["tag"].casefold(), []).append(row["line_number"])
    for lines in serial_lines.values():
        if len(lines) > 1:
            errors.append(f"Serial Number повторяется в строках: {', '.join(map(str, lines))}")
    for lines in tag_lines.values():
        if len(lines) > 1:
            errors.append(f"Tag повторяется в строках: {', '.join(map(str, lines))}")
    if not parsed:
        errors.append("Список не содержит ни одной корректной строки")
    return parsed, errors


def _sync_request_to_csv(payload: IpadInventorySyncRequest) -> str:
    if payload.csv_text.strip():
        return payload.csv_text
    if not payload.file_base64 or not payload.file_name:
        raise HTTPException(status_code=422, detail="Передайте CSV-текст или Excel-файл")
    if not payload.file_name.lower().endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="Поддерживаются только файлы .csv и .xlsx")
    try:
        content = base64.b64decode(payload.file_base64, validate=True)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook[workbook.sheetnames[0]]
        output = StringIO()
        writer = csv.writer(output)
        for row in worksheet.iter_rows(values_only=True):
            writer.writerow(["" if value is None else value for value in row])
        workbook.close()
        return output.getvalue()
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Не удалось прочитать Excel-файл") from exc


def _build_ipad_sync_plan(db: Session, csv_text: str, lock_devices: bool = False) -> tuple[dict, list[dict]]:
    rows, blocking_errors = _parse_ipad_sync_csv(csv_text)
    device_query = db.query(IpadDevice).order_by(IpadDevice.id)
    if lock_devices:
        device_query = device_query.with_for_update()
    devices = device_query.all()
    device_ids = [item.id for item in devices]
    history_ids = _history_device_ids(db, device_ids)
    active_assignments = db.query(IpadStudentAssignment).filter(
        IpadStudentAssignment.ipad_device_id.in_(device_ids),
        IpadStudentAssignment.status.in_(ACTIVE_IPAD_ASSIGNMENT_STATUSES),
    ).all() if device_ids else []
    active_by_device = {item.ipad_device_id: item for item in active_assignments}

    devices_by_serial: dict[str, IpadDevice] = {}
    duplicate_db_serials: set[str] = set()
    for device in devices:
        key = device.serial_number.strip().casefold()
        if key in devices_by_serial:
            duplicate_db_serials.add(device.serial_number)
        else:
            devices_by_serial[key] = device
    if duplicate_db_serials:
        blocking_errors.append("В базе повторяются Serial Number: " + ", ".join(sorted(duplicate_db_serials)))

    input_serials = {row["serial_number"].casefold() for row in rows}
    missing_devices = [
        device for device in devices
        if not device.is_archived and device.serial_number.strip().casefold() not in input_serials
    ]
    deletable_ids = {
        device.id for device in missing_devices
        if device.status == "AVAILABLE"
        and device.id not in active_by_device
        and device.id not in history_ids
    }
    active_device_statuses = set(ACTIVE_IPAD_ASSIGNMENT_STATUSES)
    operations: list[dict] = []
    result_rows: list[dict] = []

    for row in rows:
        existing = devices_by_serial.get(row["serial_number"].casefold())
        active_assignment = active_by_device.get(existing.id) if existing else None
        action = "UNCHANGED"
        message = "Уже совпадает с реестром"

        conflicting_tag_devices = [
            device for device in devices
            if device.tag.strip().casefold() == row["tag"].casefold()
            and (not existing or device.id != existing.id)
            and not device.is_archived
            and device.id not in deletable_ids
        ]
        if conflicting_tag_devices:
            action = "CONFLICT"
            message = "Tag уже используется другим сохраняемым iPad"
        elif existing and active_assignment:
            action = "KEEP_ACTIVE"
            message = "iPad находится в активном акте; данные защищены"
        elif existing and existing.status in active_device_statuses:
            action = "CONFLICT"
            message = "Активный статус iPad не подтверждён назначением"
        elif existing:
            changed = (
                existing.tag != row["tag"]
                or existing.model != row["model"]
                or existing.device_name != "iPad"
            )
            if existing.is_archived:
                action = "REACTIVATE"
                message = "Архивный iPad будет восстановлен"
                operations.append({"action": action, "device": existing, "row": row})
            elif changed:
                action = "UPDATE"
                message = "Tag и модель iPad будут обновлены"
                operations.append({"action": action, "device": existing, "row": row})
        else:
            action = "ADD"
            message = "Новый iPad будет добавлен как свободный"
            operations.append({"action": action, "device": None, "row": row})

        result_rows.append({**row, "action": action, "message": message, "device_id": str(existing.id) if existing else None})

    removals = []
    for device in missing_devices:
        if device.id in active_by_device or device.status in active_device_statuses:
            action = "KEEP_ACTIVE"
            message = "Отсутствует в CSV, но защищён активным актом"
        elif device.id in history_ids or device.status != "AVAILABLE":
            action = "ARCHIVE"
            message = "Отсутствует в CSV; будет скрыт с сохранением истории"
            operations.append({"action": action, "device": device, "row": None})
        else:
            action = "DELETE"
            message = "Свободный iPad без истории будет удалён"
            operations.append({"action": action, "device": device, "row": None})
        removals.append({
            "device_id": str(device.id),
            "tag": device.tag,
            "serial_number": device.serial_number,
            "action": action,
            "message": message,
        })

    action_counts = {
        action: sum(1 for item in result_rows + removals if item["action"] == action)
        for action in ("ADD", "UPDATE", "REACTIVATE", "DELETE", "ARCHIVE", "KEEP_ACTIVE", "CONFLICT", "UNCHANGED")
    }
    response = {
        "summary": {"input_rows": len(rows), **{key.lower(): value for key, value in action_counts.items()}},
        "rows": result_rows,
        "missing_devices": removals,
        "blocking_errors": blocking_errors,
        "can_apply": not blocking_errors,
        "safe_changes": len(operations),
    }
    return response, operations


@router.get("/available")
def available_ipads(
    search: Optional[str] = None,
    db: Session = Depends(get_db),
    _current_user: User = Depends(get_current_guest_or_admin_user),
):
    query = db.query(IpadDevice).filter(
        IpadDevice.status == "AVAILABLE",
        IpadDevice.is_archived.is_(False),
    )
    if search:
        value = f"%{search}%"
        query = query.filter(
            IpadDevice.tag.ilike(value)
            | IpadDevice.serial_number.ilike(value)
            | IpadDevice.model.ilike(value)
        )
    devices = query.order_by(IpadDevice.model, IpadDevice.tag, IpadDevice.serial_number).limit(500).all()
    tag_counts = dict(db.query(IpadDevice.tag, func.count(IpadDevice.id)).filter(
        IpadDevice.tag.in_([item.tag for item in devices])
    ).group_by(IpadDevice.tag).all()) if devices else {}
    return [_serialize(item, duplicate_tag_count=tag_counts.get(item.tag, 1)) for item in devices]


@router.post("/available/resolve")
def resolve_available_ipads(
    payload: IpadAvailableResolveRequest,
    db: Session = Depends(get_db),
    _current_user: User = Depends(get_current_guest_or_admin_user),
):
    requested = [str(item).strip() for item in payload.serial_numbers]
    if any(not item for item in requested):
        raise HTTPException(status_code=422, detail="Serial Number не может быть пустым")

    normalized = {item.casefold() for item in requested}
    devices = db.query(IpadDevice).filter(
        func.lower(IpadDevice.serial_number).in_(normalized),
        IpadDevice.is_archived.is_(False),
    ).order_by(IpadDevice.serial_number).all()
    devices_by_serial = {item.serial_number.casefold(): item for item in devices}

    items = []
    for serial_number in requested:
        device = devices_by_serial.get(serial_number.casefold())
        if device is None:
            items.append({
                "serial_number": serial_number,
                "match_status": "NOT_FOUND",
                "device": None,
            })
            continue
        items.append({
            "serial_number": serial_number,
            "match_status": "AVAILABLE" if device.status == "AVAILABLE" else "UNAVAILABLE",
            "device": _serialize(device),
        })
    return {"items": items}


@router.get("/groups")
def ipad_groups(
    status_value: Optional[str] = Query(None, alias="status"),
    include_archived: bool = Query(False),
    db: Session = Depends(get_db),
    _current_user: User = Depends(get_current_admin_user),
):
    query = db.query(IpadDevice.device_name, IpadDevice.model, func.count(IpadDevice.id))
    if not include_archived:
        query = query.filter(IpadDevice.is_archived.is_(False))
    if status_value:
        query = query.filter(IpadDevice.status == status_value)
    return [{"device_name": name, "model": model or "", "count": count} for name, model, count in query.group_by(IpadDevice.device_name, IpadDevice.model).order_by(func.count(IpadDevice.id).desc()).all()]


@router.get("")
def list_ipads(
    search: Optional[str] = None,
    status_value: Optional[str] = Query(None, alias="status"),
    model: Optional[str] = None,
    duplicate_tags_only: bool = Query(False),
    tag_order: Literal["asc", "desc"] = Query("asc"),
    include_archived: bool = Query(False),
    page: int = Query(1, ge=1),
    page_size: int = Query(24, ge=1, le=100),
    db: Session = Depends(get_db),
    _current_user: User = Depends(get_current_admin_user),
):
    query = db.query(IpadDevice)
    if not include_archived:
        query = query.filter(IpadDevice.is_archived.is_(False))
    if status_value:
        query = query.filter(IpadDevice.status == status_value)
    if model is not None:
        query = query.filter(IpadDevice.model == model) if model else query.filter(IpadDevice.model.is_(None))
    if search:
        value = f"%{search}%"
        query = query.filter(IpadDevice.device_name.ilike(value) | IpadDevice.model.ilike(value) | IpadDevice.tag.ilike(value) | IpadDevice.serial_number.ilike(value))
    if duplicate_tags_only:
        duplicate_tags = db.query(IpadDevice.tag).group_by(IpadDevice.tag).having(
            func.count(IpadDevice.id) > 1
        )
        query = query.filter(IpadDevice.tag.in_(duplicate_tags))
    total = query.count()
    tag_is_numeric = IpadDevice.tag.op("~")(r"^[0-9]+$")
    numeric_bucket = case((tag_is_numeric, 0), else_=1)
    tag_letters = func.lower(func.regexp_replace(IpadDevice.tag, r"[0-9]+", "", "g"))
    tag_number = cast(
        func.nullif(func.regexp_replace(IpadDevice.tag, r"[^0-9]+", "", "g"), ""),
        BigInteger,
    )
    if tag_order == "desc":
        ordering = (
            numeric_bucket.asc(),
            tag_letters.desc(),
            tag_number.desc().nullslast(),
            func.lower(IpadDevice.tag).desc(),
            IpadDevice.created_at.desc(),
        )
    else:
        ordering = (
            numeric_bucket.asc(),
            tag_letters.asc(),
            tag_number.asc().nullslast(),
            func.lower(IpadDevice.tag).asc(),
            IpadDevice.created_at.desc(),
        )
    devices = query.order_by(*ordering).offset((page - 1) * page_size).limit(page_size).all()
    device_ids = [item.id for item in devices]
    assignments = db.query(IpadStudentAssignment).filter(
        IpadStudentAssignment.ipad_device_id.in_(device_ids),
        IpadStudentAssignment.status.in_(ACTIVE_IPAD_ASSIGNMENT_STATUSES),
    ).all() if device_ids else []
    assignment_by_device = {item.ipad_device_id: item for item in assignments}
    history_device_ids = _history_device_ids(db, device_ids)
    tag_counts = dict(db.query(IpadDevice.tag, func.count(IpadDevice.id)).filter(
        IpadDevice.tag.in_([item.tag for item in devices])
    ).group_by(IpadDevice.tag).all()) if devices else {}
    return {"items": [
        _serialize(
            item,
            assignment_by_device.get(item.id),
            tag_counts.get(item.tag, 1),
            has_history=item.id in history_device_ids,
        )
        for item in devices
    ], "total": total, "page": page, "page_size": page_size}


@router.post("/sync/preview")
def preview_ipad_inventory_sync(
    payload: IpadInventorySyncRequest,
    db: Session = Depends(get_db),
    _current_user: User = Depends(get_current_admin_user),
):
    preview, _operations = _build_ipad_sync_plan(db, _sync_request_to_csv(payload))
    return preview


@router.post("/sync/apply")
def apply_ipad_inventory_sync(
    payload: IpadInventorySyncRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_admin_user),
):
    preview, operations = _build_ipad_sync_plan(db, _sync_request_to_csv(payload), lock_devices=True)
    if preview["blocking_errors"]:
        raise HTTPException(status_code=422, detail=preview["blocking_errors"])

    for operation in operations:
        action = operation["action"]
        device = operation["device"]
        row = operation["row"]
        if action == "DELETE":
            record_audit(db, current_user, "IPAD_DEVICE", device.id, "IPAD_DEVICE_SYNC_DELETED", {
                "tag": device.tag,
                "serial_number": device.serial_number,
            })
            db.delete(device)
        elif action == "ARCHIVE":
            device.is_archived = True
            record_audit(db, current_user, "IPAD_DEVICE", device.id, "IPAD_DEVICE_SYNC_ARCHIVED", {
                "tag": device.tag,
                "serial_number": device.serial_number,
            })
        elif action in {"UPDATE", "REACTIVATE"}:
            device.tag = row["tag"]
            device.model = row["model"]
            device.device_name = row["device_name"]
            device.is_archived = False
            record_audit(db, current_user, "IPAD_DEVICE", device.id, f"IPAD_DEVICE_SYNC_{action}", {
                "tag": device.tag,
                "serial_number": device.serial_number,
            })
        elif action == "ADD":
            device = IpadDevice(
                id=uuid.uuid4(),
                device_name=row["device_name"],
                model=row["model"],
                tag=row["tag"],
                serial_number=row["serial_number"],
                status="AVAILABLE",
                is_archived=False,
            )
            db.add(device)
            record_audit(db, current_user, "IPAD_DEVICE", device.id, "IPAD_DEVICE_SYNC_ADDED", {
                "tag": device.tag,
                "serial_number": device.serial_number,
            })

    record_audit(db, current_user, "IPAD_INVENTORY", current_user.id, "IPAD_INVENTORY_SYNC_APPLIED", preview["summary"])
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Синхронизация конфликтует с текущим реестром iPad")
    return {**preview, "applied": True}


@router.post("", status_code=status.HTTP_201_CREATED)
def create_ipad(data: IpadDeviceCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_admin_user)):
    values = data.model_dump()
    values.update(device_name=data.device_name.strip(), model=data.model.strip() if data.model else None, tag=data.tag.strip(), serial_number=data.serial_number.strip())
    if not values["device_name"] or not values["tag"] or not values["serial_number"]:
        raise HTTPException(status_code=422, detail="Device name, Tag и Serial Number обязательны")
    device = IpadDevice(**values)
    db.add(device)
    db.flush()
    record_audit(db, current_user, "IPAD_DEVICE", device.id, "IPAD_DEVICE_CREATED", {"tag": device.tag})
    try: db.commit()
    except IntegrityError:
        db.rollback(); raise HTTPException(status_code=409, detail="Serial Number уже используется")
    db.refresh(device)
    return _serialize(device)


@router.post("/bulk", status_code=status.HTTP_201_CREATED)
def bulk_create_ipads(data: IpadDeviceBulkCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_admin_user)):
    name = data.device_name.strip()
    model = data.model.strip() if data.model else None
    rows = [(item.tag.strip(), item.serial_number.strip()) for item in data.devices]
    if not name or not rows or any(not tag or not serial for tag, serial in rows):
        raise HTTPException(status_code=422, detail="Заполните общие поля, Tag и Serial Number")
    tags = [item[0] for item in rows]; serials = [item[1] for item in rows]
    if len(serials) != len(set(serials)):
        raise HTTPException(status_code=422, detail="Serial Number повторяется в списке")
    conflicts = db.query(IpadDevice).filter(IpadDevice.serial_number.in_(serials)).all()
    if conflicts: raise HTTPException(status_code=409, detail="Serial Number уже существуют: " + ", ".join(item.serial_number for item in conflicts))
    created = []
    for tag, serial in rows:
        device = IpadDevice(id=uuid.uuid4(), device_name=name, model=model, tag=tag, serial_number=serial, status=data.status)
        db.add(device); created.append(device)
        record_audit(db, current_user, "IPAD_DEVICE", device.id, "IPAD_DEVICE_CREATED", {"source": "bulk", "tag": tag})
    try: db.commit()
    except IntegrityError:
        db.rollback(); raise HTTPException(status_code=409, detail="Serial Number уже используется")
    return {"created": len(created)}


@router.patch("/{device_id}")
def update_ipad(device_id: UUID, data: IpadDeviceUpdate, db: Session = Depends(get_db), current_user: User = Depends(get_current_admin_user)):
    device = db.query(IpadDevice).filter(IpadDevice.id == device_id).first()
    if not device: raise HTTPException(status_code=404, detail="iPad не найден")
    active = db.query(IpadStudentAssignment.id).filter(IpadStudentAssignment.ipad_device_id == device.id, IpadStudentAssignment.status.in_(ACTIVE_IPAD_ASSIGNMENT_STATUSES)).first()
    updates = data.model_dump(exclude_unset=True)
    if active and any(key in updates for key in ("tag", "serial_number", "status")):
        raise HTTPException(status_code=409, detail="Активный iPad управляется через акт")
    for key, value in updates.items(): setattr(device, key, value.strip() if isinstance(value, str) else value)
    record_audit(db, current_user, "IPAD_DEVICE", device.id, "IPAD_DEVICE_UPDATED", {"fields": list(updates)})
    try: db.commit()
    except IntegrityError:
        db.rollback(); raise HTTPException(status_code=409, detail="Serial Number уже используется")
    db.refresh(device); return _serialize(device)


@router.get("/{device_id}/history")
def ipad_history(
    device_id: UUID,
    db: Session = Depends(get_db),
    _current_user: User = Depends(get_current_admin_user),
):
    """Паспорт устройства: все выдачи, повреждения, замены и ремонты iPad."""
    device = db.query(IpadDevice).filter(IpadDevice.id == device_id).first()
    if not device:
        raise HTTPException(status_code=404, detail="iPad не найден")

    events: list[dict] = []

    # Выдачи: каждое назначение ученику через акт.
    assignments = db.query(IpadStudentAssignment).filter(
        IpadStudentAssignment.ipad_device_id == device.id
    ).all()
    act_ids = {item.act_id for item in assignments}
    acts = {act.id: act for act in db.query(Act).filter(Act.id.in_(act_ids)).all()} if act_ids else {}
    for item in assignments:
        act = acts.get(item.act_id)
        events.append({
            "type": "ASSIGNMENT",
            "title": f"Выдан ученику: {item.student_name}",
            "detail": f"Акт ACT-{str(item.act_id).split('-')[0].upper()}" + (f" · {act.item_name}" if act else ""),
            "status": item.status,
            "act_id": str(item.act_id),
            "created_at": item.assigned_at.isoformat(),
        })

    # Приложения: замены (старый/новый), выбытия, поздние возвраты — по payload.
    device_id_str = str(device.id)
    appendices = db.query(IpadActAppendix).filter(IpadActAppendix.status == "APPLIED").all()
    for appendix in appendices:
        payload = appendix.payload_json or {}
        related = device_id_str in {
            str(payload.get("old_device_id")), str(payload.get("new_device_id")), str(payload.get("device_id")),
        }
        if not related and appendix.operation_type == "YEAR_END_RETURN":
            related = any(str(item.get("device_id")) == device_id_str for item in payload.get("items", []) if isinstance(item, dict))
        if not related:
            continue
        titles = {
            "IPAD_REPLACEMENT": "Замена iPad",
            "STUDENT_DEPARTURE": "Выбытие ученика",
            "STUDENT_ADDITION": "Добавление ученика",
            "LATE_RETURN": "Поздний возврат",
            "YEAR_END_RETURN": "Годовой возврат",
        }
        role = ""
        if appendix.operation_type == "IPAD_REPLACEMENT":
            role = " (заменён)" if str(payload.get("old_device_id")) == device_id_str else " (выдан взамен)"
        condition = payload.get("reason_label") or payload.get("return_condition_label") or payload.get("condition_label")
        events.append({
            "type": "APPENDIX",
            "title": titles.get(appendix.operation_type, appendix.operation_type) + role,
            "detail": (payload.get("student_name") or "") + (f" · {condition}" if condition else ""),
            "status": appendix.operation_type,
            "act_id": str(appendix.act_id),
            "created_at": (appendix.applied_at or appendix.created_at).isoformat(),
        })

    events.sort(key=lambda item: item["created_at"], reverse=True)
    return {
        "device": _serialize(device, has_history=bool(events)),
        "events": events,
    }


@router.delete("/{device_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_ipad(device_id: UUID, db: Session = Depends(get_db), current_user: User = Depends(get_current_admin_user)):
    device = db.query(IpadDevice).filter(IpadDevice.id == device_id).with_for_update().first()
    if not device: raise HTTPException(status_code=404, detail="iPad не найден")
    if device.status in {"RESERVED", "ISSUED", "RETURN_PENDING"}:
        raise HTTPException(status_code=409, detail="Активный или зарезервированный iPad нельзя удалить")
    if device.id in _history_device_ids(db, [device.id]):
        raise HTTPException(status_code=409, detail="iPad имеет историю актов и не может быть удалён")
    record_audit(db, current_user, "IPAD_DEVICE", device.id, "IPAD_DEVICE_DELETED", {"tag": device.tag}); db.flush(); db.delete(device); db.commit()
